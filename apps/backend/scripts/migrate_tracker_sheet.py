"""Load the Job Priority workbook into Postgres.

The sheet is the system of record until cutover, so this is re-runnable: every row
is upserted by ``job_id`` and running it twice changes nothing. That matters because
the real cutover is "migrate, keep using the sheet for a day, migrate the delta".

Usage::

    python -m scripts.migrate_tracker_sheet --xlsx "/path/Linkedin JD update.xlsx"

Reads the sheet tabs, maps them onto the tracker tables, recomputes ``sort_key`` and
the Assigned-queue flags, and reports what landed. Apify tokens are written but never
printed.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime
from typing import Any, Iterator

import openpyxl
from sqlalchemy import create_engine, func, select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.tracker.enums import is_assignee, is_dead_status  # noqa: E402
from app.tracker.models import (  # noqa: E402
    TrackerApifyAccount,
    TrackerJob,
    TrackerJobDoc,
    TrackerSetting,
    init_tracker_schema,
)
from app.tracker.ranking import build_sort_key  # noqa: E402

#: Header rows differ per tab because each has its own banner above the header.
HEADER_ROW = {"Job_Priority": 6, "Assigned": 3, "Raw_Data": 1}

RAW_REF_LIMIT = 40_000
JD_TEXT_LIMIT = 12_000


def rows_of(wb: openpyxl.Workbook, sheet: str) -> Iterator[dict[str, Any]]:
    """Yield each data row of a tab as a dict keyed by its header."""
    ws = wb[sheet]
    header_at = HEADER_ROW[sheet]
    header: list[str] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < header_at:
            continue
        if i == header_at:
            header = [str(c).strip() if c else f"col{n}" for n, c in enumerate(row)]
            continue
        if row is None or all(c in (None, "") for c in row):
            continue
        yield dict(zip(header, row))


def text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def as_int(v: Any) -> int | None:
    if v is None or v == "":
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def as_bool(v: Any) -> bool | None:
    s = text(v).lower()
    if s in ("true", "yes", "1"):
        return True
    if s in ("false", "no", "0"):
        return False
    return None


def as_dt(v: Any) -> datetime | None:
    if isinstance(v, datetime):
        return v
    s = text(v)
    if not s:
        return None
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None


def company_key(name: Any) -> str:
    """Normalised company name, matching how the sheet groups for de-duplication."""
    return re.sub(r"[^a-z0-9]+", " ", text(name).lower()).strip()


def split_ids(v: Any) -> list[str]:
    return [p for p in re.split(r"[,\s]+", text(v)) if p]


def extract_jd(raw_ref: str) -> str:
    """Pull the description out of a stored payload, best-effort.

    The payload is JSON for rows the current scraper wrote, and plain text for older
    ones, so an unparseable value falls back to the raw string rather than failing
    the migration.
    """
    if not raw_ref:
        return ""
    try:
        obj = json.loads(raw_ref)
    except (ValueError, TypeError):
        return raw_ref[:JD_TEXT_LIMIT]
    if isinstance(obj, dict):
        for key in ("description", "jobDescription", "job_description", "descriptionText"):
            if isinstance(obj.get(key), str) and obj[key].strip():
                return obj[key][:JD_TEXT_LIMIT]
        info = obj.get("job_info") or obj.get("jobInfo") or {}
        if isinstance(info, dict) and isinstance(info.get("description"), str):
            return info["description"][:JD_TEXT_LIMIT]
    return raw_ref[:JD_TEXT_LIMIT]


def upsert(session: Session, model: Any, rows: list[dict[str, Any]], pk: list[str]) -> int:
    """Insert or update in batches, so a re-run is a no-op rather than a duplicate."""
    if not rows:
        return 0
    written = 0
    for start in range(0, len(rows), 400):
        chunk = rows[start : start + 400]
        stmt = insert(model).values(chunk)
        update = {c: stmt.excluded[c] for c in chunk[0] if c not in pk}
        session.execute(stmt.on_conflict_do_update(index_elements=pk, set_=update))
        written += len(chunk)
    session.commit()
    return written


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--database-url", default=os.environ.get(
        "TRACKER_DATABASE_URL", "postgresql+psycopg://postgres:postgres@localhost:5432/som_career_coach"))
    ap.add_argument("--user-id", default="owner")
    ap.add_argument("--raw-limit", type=int, default=0,
                    help="cap Raw_Data rows (0 = all); the payloads are large")
    args = ap.parse_args()

    engine = create_engine(args.database_url, future=True)
    init_tracker_schema(engine)
    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)

    # ---- Assigned first: its columns belong on the same row as the job -------
    assigned: dict[str, dict[str, Any]] = {}
    for r in rows_of(wb, "Assigned"):
        jid = text(r.get("job_id"))
        if jid:
            assigned[jid] = r
    print(f"Assigned      : {len(assigned)} rows to fold in")

    # ---- Job_Priority -------------------------------------------------------
    jobs: list[dict[str, Any]] = []
    seen: set[str] = set()
    skipped_no_id = 0
    for r in rows_of(wb, "Job_Priority"):
        jid = text(r.get("job_id"))
        if not jid:
            skipped_no_id += 1
            continue
        if jid in seen:  # the sheet can hold a stale duplicate; last one wins
            jobs = [j for j in jobs if j["job_id"] != jid]
        seen.add(jid)

        a = assigned.get(jid, {})
        status, owner = text(r.get("status")) or "New", text(r.get("owner"))
        posted, imported = as_dt(r.get("posted")), as_dt(r.get("imported_at"))
        score = as_int(r.get("score"))
        sort_key = build_sort_key(status, r.get("priority"), posted, imported, score)
        in_queue = is_assignee(owner) and not is_dead_status(status)

        jobs.append({
            "job_id": jid, "user_id": args.user_id,
            "priority": text(r.get("priority")) or None, "score": score,
            "us_visa": text(r.get("us_visa")) or None,
            "us_visa_reason": text(r.get("us_visa_reason")) or None,
            "summary": text(r.get("summary")) or None, "why": text(r.get("why")) or None,
            "level_normalized": text(r.get("level_normalized")) or None,
            "requires_people_mgmt": as_bool(r.get("requires_people_mgmt")),
            "required_yoe_pm": as_int(r.get("required_yoe_pm")),
            "required_yoe_total": as_int(r.get("required_yoe_total")),
            "scored_at": as_dt(r.get("scored_at")),
            "scoring_fingerprint": text(r.get("scoring_fingerprint")) or None,
            "company": text(r.get("company")) or None, "title": text(r.get("title")) or None,
            "title_level": text(r.get("title_level")) or None,
            "jd_implied_level": text(r.get("jd_implied_level")) or None,
            "location": text(r.get("location")) or None,
            "posted": posted, "applicants": text(r.get("applicants")) or None,
            "imported_at": imported,
            "status": status, "owner": owner, "action": text(r.get("action")),
            "referral_contact": text(r.get("referral_contact")) or None,
            "notes": text(a.get("notes")) or None,
            "assigned_updated_at": as_dt(a.get("updated_at")),
            "applied_at": as_dt(a.get("applied_at")),
            "merged_job_ids": split_ids(r.get("merged_job_ids")) or [jid],
            "jd_fingerprint": text(r.get("jd_fingerprint")) or None,
            "company_key": company_key(r.get("company")) or None,
            "dead_links": {}, "resume_file": None,
            "sort_key": sort_key, "assigned_sort_key": sort_key,
            "in_assigned_queue": in_queue,
        })

    written = upsert(Session(engine), TrackerJob, jobs, ["job_id"])
    print(f"Job_Priority  : {written} jobs ({skipped_no_id} rows had no job_id)")

    # ---- Raw_Data -----------------------------------------------------------
    # Keyed by job_id: once delete markers are stripped, a tombstoned row and a live
    # row can share an id, and Postgres rejects two conflicting rows in one upsert.
    by_id: dict[str, dict[str, Any]] = {}
    orphans = 0
    tombstoned = 0
    for r in rows_of(wb, "Raw_Data"):
        raw_id = text(r.get("job_id"))
        if not raw_id:
            continue
        # A leading '~' is the sheet's delete marker: prune renames the id so the row
        # sorts below every real one, then deletes the block. A run that hit the
        # 6-minute limit mid-sweep leaves the marker behind. The id underneath is
        # still good, so recover it and let the new prune decide the row's fate.
        jid = raw_id.lstrip("~")
        if jid != raw_id:
            tombstoned += 1
        if jid not in seen:  # payload for a job that no longer exists
            orphans += 1
            continue
        raw = text(r.get("raw_ref"))[:RAW_REF_LIMIT]
        candidate = {
            "job_id": jid, "raw_ref": raw or None, "jd_text": extract_jd(raw) or None,
            "posted_at": as_dt(r.get("posted_at")), "listed_at": as_dt(r.get("listed_at")),
            "imported_at": as_dt(r.get("imported_at")),
        }
        existing = by_id.get(jid)
        # Keep whichever copy carries more payload; a tombstoned row is often the
        # emptied one, and the live row is what the app should score against.
        if existing is None or len(raw) > len(existing["raw_ref"] or ""):
            by_id[jid] = candidate
        if args.raw_limit and len(by_id) >= args.raw_limit:
            break

    docs = list(by_id.values())
    written = upsert(Session(engine), TrackerJobDoc, docs, ["job_id"])
    print(f"Raw_Data      : {written} payloads ({orphans} orphans skipped, "
          f"{tombstoned} carried a leftover delete marker)")

    # ---- Settings + Prompt --------------------------------------------------
    settings: list[dict[str, Any]] = []
    ws = wb["Settings"]
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        k = text(row[0])
        if k.startswith("#") or k.lower() == "setting_key":
            continue
        settings.append({"key": k, "user_id": args.user_id,
                         "value": text(row[1]) if len(row) > 1 else ""})

    prompt_rows = list(wb["Prompt"].iter_rows(values_only=True))
    def cell(n: int) -> str:
        return text(prompt_rows[n][1]) if len(prompt_rows) > n and len(prompt_rows[n]) > 1 else ""
    settings.append({"key": "TARGET_PROFILE", "user_id": args.user_id, "value": cell(1)})
    settings.append({"key": "SCORING_INSTRUCTIONS", "user_id": args.user_id, "value": cell(2)})

    written = upsert(Session(engine), TrackerSetting, settings, ["key"])
    secret = sum(1 for s in settings if "KEY" in s["key"].upper() and s["value"])
    print(f"Settings      : {written} keys ({secret} hold secrets, not printed)")

    # ---- Apify accounts -----------------------------------------------------
    accounts: list[dict[str, Any]] = []
    for pos, row in enumerate(wb["Apify_Accounts"].iter_rows(values_only=True)):
        if not row or not row[0]:
            continue
        batch = text(row[1]) if len(row) > 1 else ""
        if not batch.startswith("apify_"):  # header and instruction rows
            continue
        accounts.append({
            "label": text(row[0]), "user_id": args.user_id, "batch_key": batch,
            "detail_key": text(row[2]) if len(row) > 2 else "",
            "active": as_bool(row[3] if len(row) > 3 else None) is not False,
            "position": pos,
        })
    written = upsert(Session(engine), TrackerApifyAccount, accounts, ["label"])
    print(f"Apify accounts: {written} ({[a['label'] for a in accounts]})")

    # ---- what landed --------------------------------------------------------
    with Session(engine) as s:
        total = s.scalar(select(func.count()).select_from(TrackerJob))
        queued = s.scalar(
            select(func.count()).select_from(TrackerJob).where(TrackerJob.in_assigned_queue.is_(True)))
        by_status = s.execute(
            select(TrackerJob.status, func.count()).group_by(TrackerJob.status)
            .order_by(func.count().desc())).all()
    print(f"\nIn database   : {total} jobs, {queued} in the assigned queue")
    print("By status     : " + ", ".join(f"{st or '(blank)'}={n}" for st, n in by_status))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
